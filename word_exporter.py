from disassembly_loader import build_guide
from docx import Document 
from docx.shared import Inches
import os
import sys

# Input argument check
if len(sys.argv) < 2:
    print("Error: You have to pass the path of the json file you want to disassemble")
    print("Usage: python3 test.py <filename>.json")
    sys.exit(1) 

json_path = sys.argv[1]

try:
    # Loading of the guide
    guide = build_guide(json_path, include_bom=True)

    # Word document creation
    doc = Document()
    doc.add_heading(f"Disassembly guide: {guide.product.name}", 0)

    # [Readme Implementation] Print validation warnings at the beginning if present
    if guide.warnings:
        doc.add_heading("Validation warnings", level=1)
        for warning in guide.warnings:
            p_warn = doc.add_paragraph(style='Normal')
            r_warn = p_warn.add_run(f"[{warning.severity.upper()}] Rule '{warning.rule}': {warning.message}")
            r_warn.italic = True
    
    # Printing bill of materials
    if hasattr(guide, 'bill_of_materials') and guide.bill_of_materials:
        doc.add_heading("Bill of materials", level=1)
        
        # Table with 1 row heading and 4 cols
        table = doc.add_table(rows=1, cols=4)
        table.style = 'Table Grid'
        
        # column headings
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'Component'
        hdr_cells[1].text = 'Weight'
        hdr_cells[2].text = 'Material'
        hdr_cells[3].text = 'Color'
        
        # Add the bold to the text of the headings
        for cell in hdr_cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True
        
        # Fill the table
        for item in guide.bill_of_materials:
            row_cells = table.add_row().cells
            
            # component name
            row_cells[0].text = str(item.name)
            
            # weight
            if item.weight is not None:
                row_cells[1].text = f"{int(item.weight) if isinstance(item.weight, float) and item.weight.is_integer() else item.weight} g"
            else:
                row_cells[1].text = "—"
                
            # material
            row_cells[2].text = str(item.material) if item.material is not None else "—"
            
            # color
            row_cells[3].text = str(item.color) if item.color is not None else "—"
            
        col_widths = [Inches(3.0), Inches(1.0), Inches(1.0), Inches(1.0)]
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = width
                
        doc.add_paragraph()
    
    # Pre-collect all parent source names that will be disassembled throughout the guide
    active_parent_names = set()
    for step in guide.steps:
        # Replicate the logic to determine where a step is dismounting from
        source_name = guide.product.name if step.index == 1 or not guide.steps[step.index - 2].continues_as else guide.steps[step.index - 2].continues_as.name
        active_parent_names.add(source_name)
    
    # Iterating through steps (and writing them on the final docx file)
    for step in guide.steps:
        
        # Printing the composite sub-object where you are dismounting from
        current_source = guide.product.name if step.index == 1 or not guide.steps[step.index - 2].continues_as else guide.steps[step.index - 2].continues_as.name
        doc.add_paragraph() 
        doc.add_heading(f"Disassembling from {current_source}", level=1)
        
        # Printing action
        p_action = doc.add_paragraph()
        r_action = p_action.add_run(f"Action {step.index}: {step.operation}")
        r_action.bold = True
        
        # [Readme Implementation] Print required tools for this specific step if any are needed
        if step.tools_required:
            p_tools = doc.add_paragraph()
            r_tools_label = p_tools.add_run("Tools required: ")
            r_tools_label.bold = True
            p_tools.add_run(", ".join(step.tools_required))
        
        # Steps printed as bulleted list
        if step.actions:
            for action in step.actions:
                doc.add_paragraph(action.text, style='List Bullet 2')

        # Printing removed pieces
        # I include only "sub-roots (="composite") that will be disassembled later"
        collected_sub_roots = []
        
        # Add the main continuation component if it gets worked on downstream
        if step.continues_as and step.continues_as.name in active_parent_names:
            collected_sub_roots.append(step.continues_as.name)
            
        # Add any other isolated parallel branches from outputs that get worked on downstream
        for out_component in step.outputs:
            if out_component.name in active_parent_names and out_component.name not in collected_sub_roots:
                collected_sub_roots.append(out_component.name)

        if collected_sub_roots:
            p = doc.add_paragraph()
            # Loop over each sub-root to build the exact sequence of resumes
            resume_phrases = [f'"Disassembling from {name}"' for name in collected_sub_roots]
            
            r = p.add_run("Then resume from " + ", ".join(resume_phrases))
            r.bold = True

    # File saving
    base_name = os.path.splitext(os.path.basename(json_path))[0]
    doc.save(f"{base_name}.docx")
    
    print(f"Word (.docx) file created: {os.path.abspath(f'{base_name}.docx')}")

except Exception as e:
    print(f"Error: {e}")