import os
import sys
from docx import Document
from disassembly_loader import build_guide


def export_to_word(json_path):
    """
    Reads the disassembly JSON file via Rubin's loader
    and exports a formatted Word Document (.docx).
    """
    # Initialize the guide using the shared loader logic
    guide = build_guide(json_path)
    doc = Document()

    # Main Title
    doc.add_heading(f"Disassembly guide: {guide.product.name}", 0)

    # 1. Validation Warnings Section (Simone's logic integration)
    if guide.warnings:
        doc.add_heading("Validation warnings", level=1)
        for warning in guide.warnings:
            p_warn = doc.add_paragraph(style='Normal')
            r_warn = p_warn.add_run(f"[{warning.severity.upper()}] Rule '{warning.rule}': {warning.message}")
            r_warn.italic = True

    # Pre-collect active parent components for tree tracking
    active_parent_names = set()
    for step in guide.steps:
        source_name = guide.product.name if step.index == 1 or not guide.steps[step.index - 2].continues_as else \
        guide.steps[step.index - 2].continues_as.name
        active_parent_names.add(source_name)

    # 2. Iterate and write steps
    for step in guide.steps:
        current_source = guide.product.name if step.index == 1 or not guide.steps[step.index - 2].continues_as else \
        guide.steps[step.index - 2].continues_as.name
        doc.add_paragraph()
        doc.add_heading(f"Disassembling from {current_source}", level=1)

        # Primary Action Title
        p_action = doc.add_paragraph()
        r_action = p_action.add_run(f"Action {step.index}: {step.operation}")
        r_action.bold = True

        # Required Tools Badge/Text
        if step.tools_required:
            p_tools = doc.add_paragraph()
            r_tools_label = p_tools.add_run("Tools required: ")
            r_tools_label.bold = True
            p_tools.add_run(", ".join(step.tools_required))

        # Detailed Sub-actions (Bulleted List)
        if step.actions:
            for action in step.actions:
                doc.add_paragraph(action.text, style='List Bullet 2')

        # Continuation Resumes tracking
        collected_sub_roots = []
        if step.continues_as and step.continues_as.name in active_parent_names:
            collected_sub_roots.append(step.continues_as.name)
        for out_component in step.outputs:
            if out_component.name in active_parent_names and out_component.name not in collected_sub_roots:
                collected_sub_roots.append(out_component.name)

        if collected_sub_roots:
            p = doc.add_paragraph()
            resume_phrases = [f'"Disassembling from {name}"' for name in collected_sub_roots]
            r = p.add_run("Then resume from " + ", ".join(resume_phrases))
            r.bold = True

    # Save output file matching the original JSON name
    base_name = os.path.splitext(os.path.basename(json_path))[0]
    output_path = f"{base_name}.docx"
    doc.save(output_path)
    return os.path.abspath(output_path)