import os
from pptx import Presentation
from disassembly_loader import build_guide


def export_to_pptx(json_path):
    """
    Reads the disassembly JSON file via Rubin's loader
    and exports a professional PowerPoint Presentation (.pptx).
    """
    guide = build_guide(json_path)
    prs = Presentation()

    # 1. Slide 1: Title Slide (Layout 0)
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    slide.shapes.title.text = f"Disassembly Guide\n{guide.product.name}"
    slide.placeholders[1].text = "Generated via Disassembly Wizard Desktop App"

    # 2. Slide 2: Validation Warnings (Layout 1) - Only if warnings exist
    if guide.warnings:
        bullet_slide_layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(bullet_slide_layout)
        slide.shapes.title.text = "Parser Validation Warnings"
        tf = slide.placeholders[1].text_frame

        for warning in guide.warnings:
            p = tf.add_paragraph()
            p.text = f"• [{warning.severity.upper()}] Rule '{warning.rule}': {warning.message}"
            p.level = 0

    # Pre-collect active parent components for tree continuity (1:1 with Simone's logic)
    active_parent_names = set()
    for step in guide.steps:
        source_name = guide.product.name if step.index == 1 or not guide.steps[step.index - 2].continues_as else \
        guide.steps[step.index - 2].continues_as.name
        active_parent_names.add(source_name)

    # 3. Create a dedicated slide for each Disassembly Step
    for step in guide.steps:
        current_source = guide.product.name if step.index == 1 or not guide.steps[step.index - 2].continues_as else \
        guide.steps[step.index - 2].continues_as.name

        bullet_slide_layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(bullet_slide_layout)

        # Set Slide Title as the Main Action
        slide.shapes.title.text = f"Action {step.index}: {step.operation}"

        # Select the body text frame
        tf = slide.placeholders[1].text_frame
        tf.text = f"Source Location: Dismounting from '{current_source}'"

        # Inject Tools Required
        if step.tools_required:
            p_tools = tf.add_paragraph()
            p_tools.text = f"Tools Required: {', '.join(step.tools_required)}"
            p_tools.level = 0

        # Inject Detailed Sub-steps (Indented Level 1)
        if step.actions:
            for action in step.actions:
                p_act = tf.add_paragraph()
                p_act.text = action.text
                p_act.level = 1

        # Inject Continuation Guidance (Tree Tracing)
        collected_sub_roots = []
        if step.continues_as and step.continues_as.name in active_parent_names:
            collected_sub_roots.append(step.continues_as.name)
        for out_component in step.outputs:
            if out_component.name in active_parent_names and out_component.name not in collected_sub_roots:
                collected_sub_roots.append(out_component.name)

        if collected_sub_roots:
            p_next = tf.add_paragraph()
            resume_phrases = [f'"{name}"' for name in collected_sub_roots]
            p_next.text = f"Workflow Continuity: Then resume from {', '.join(resume_phrases)}"
            p_next.level = 0

    # Save output presentation file matching the JSON filename
    base_name = os.path.splitext(os.path.basename(json_path))[0]
    output_path = f"{base_name}.pptx"
    prs.save(output_path)
    return os.path.abspath(output_path)